#!/usr/bin/env python3
"""
OrphaAI — Database Seeding Script
Populates local SQLite database (orphaai.db) with Broad Institute Drug Repurposing Hub dataset records,
including action_type and therapeutic_direction annotations.

Usage:
    python scripts/seed_db.py [optional_path_to_data_file]
"""

import os
import sys
import re
import json
import csv
from pathlib import Path

# Ensure backend directory is in python path
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent if SCRIPT_DIR.name in ("scripts", "backend") else SCRIPT_DIR
BACKEND_DIR = PROJECT_ROOT / "backend" / "backend" if (PROJECT_ROOT / "backend" / "backend").exists() else PROJECT_ROOT / "backend"

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

try:
    from config import CANONICAL_DB_PATH
except ImportError:
    try:
        from backend.config import CANONICAL_DB_PATH
    except ImportError:
        CANONICAL_DB_PATH = os.path.abspath(os.path.join(PROJECT_ROOT, "instance", "orphaai.db"))

from sqlalchemy import create_engine, Column, Integer, String, Text
from sqlalchemy.orm import declarative_base, sessionmaker

Base = declarative_base()

class RepurposingRecord(Base):
    __tablename__ = "repurposing_records"

    id                  = Column(Integer, primary_key=True, autoincrement=True)
    disease             = Column(String(255), index=True, nullable=False)
    drug_name           = Column(String(255), index=True, nullable=False)
    clinical_phase      = Column(String(100), default="Investigational")
    disease_area        = Column(String(255), default="General")
    mechanism_of_action = Column(Text, default="")
    target_genes        = Column(Text, default="")
    evidence_strength   = Column(String(100), default="Investigational")
    action_type         = Column(String(50), default="Unknown", index=True)
    therapeutic_direction = Column(String(30), default="Modulation", index=True)

    def to_dict(self):
        return {
            "id": self.id,
            "disease": self.disease,
            "drug_name": self.drug_name,
            "clinical_phase": self.clinical_phase,
            "disease_area": self.disease_area,
            "mechanism_of_action": self.mechanism_of_action,
            "target_genes": self.target_genes,
            "evidence_strength": self.evidence_strength,
            "action_type": self.action_type or "Unknown",
            "therapeutic_direction": self.therapeutic_direction or "Modulation",
        }

def map_evidence_strength(clinical_phase: str, explicit_evidence: str = "") -> str:
    phase = str(clinical_phase or "").strip().lower()
    explicit = str(explicit_evidence or "").strip().lower()

    if any(k in phase for k in ["launched", "approved", "phase 4", "phase iv"]) or any(k in explicit for k in ["launched", "approved", "very strong", "high"]):
        return "Approved / High Evidence"
    elif any(k in phase for k in ["phase 3", "phase 2", "phase iii", "phase ii"]) or any(k in explicit for k in ["clinical", "moderate"]):
        return "Moderate Evidence"
    else:
        return "Investigational"

def parse_action_direction(moa_str: str) -> tuple:
    moa = str(moa_str or "").strip().lower()
    if not moa:
        return "Unknown", "Modulation"

    # Inhibitor / Antagonist / Suppressor / Blocker
    if any(w in moa for w in ["inhibitor", "antagonist", "blocker", "suppressor", "degrader", "cleaver"]):
        if "channel blocker" in moa:
            action_type = "Channel Blocker"
        elif "antagonist" in moa:
            action_type = "Antagonist"
        else:
            action_type = "Inhibitor"
        return action_type, "Downregulation/Inhibition"

    # Agonist / Activator / Stimulant / Potentiator
    elif any(w in moa for w in ["agonist", "activator", "stimulant", "potentiator", "enhancer", "inducer"]):
        if "agonist" in moa:
            action_type = "Agonist"
        else:
            action_type = "Activator"
        return action_type, "Upregulation/Activation"

    # Modulator / Binder / Chaperone
    elif any(w in moa for w in ["modulator", "binder", "chaperone", "ligand", "opener"]):
        if "allosteric" in moa:
            action_type = "Allosteric Modulator"
        else:
            action_type = "Modulator"
        return action_type, "Modulation"

    return "Unknown", "Modulation"

def explode_delimiters(val_str: str) -> list:
    if not val_str:
        return []
    parts = re.split(r"[|;,]", str(val_str))
    cleaned = [p.strip() for p in parts if p.strip()]
    return list(dict.fromkeys(cleaned))

def title_case_disease(disease_str: str) -> str:
    if not disease_str or not disease_str.strip():
        return ""
    clean = disease_str.strip()
    acronyms = {
        "COVID-19": "COVID-19", "ALS": "ALS", "HIV": "HIV", "AIDS": "AIDS",
        "NSCLC": "NSCLC", "ADHD": "ADHD", "COPD": "COPD", "CF": "CF", "PD": "PD", "MS": "MS"
    }
    if clean.upper() in acronyms:
        return acronyms[clean.upper()]

    words = clean.split()
    formatted = []
    for w in words:
        w_upper = w.upper()
        if w_upper in acronyms:
            formatted.append(acronyms[w_upper])
        elif "'" in w:
            parts = w.split("'")
            formatted.append(parts[0].capitalize() + "'" + parts[1].lower())
        elif "-" in w:
            parts = w.split("-")
            formatted.append("-".join(p.capitalize() if p.upper() not in acronyms else acronyms[p.upper()] for p in parts))
        elif len(w) <= 3 and w.isupper():
            formatted.append(w)
        else:
            formatted.append(w.capitalize())
    return " ".join(formatted)

def find_data_file(custom_path=None) -> Path:
    if custom_path and os.path.exists(custom_path):
        return Path(custom_path)

    search_candidates = [
        BACKEND_DIR / "data" / "repurposing_drugs.tsv",
        PROJECT_ROOT / "backend" / "backend" / "data" / "repurposing_drugs.tsv",
        PROJECT_ROOT / "repurposing_drugs.tsv",
        BACKEND_DIR / "data" / "repurposing_drugs.xlsx",
        BACKEND_DIR / "data" / "repurposing_dataset.json",
    ]

    for candidate in search_candidates:
        if candidate.exists():
            return candidate

    raise FileNotFoundError("Could not locate Broad Repurposing Hub data file (repurposing_drugs.tsv or .xlsx or .json).")

def parse_data_file(file_path: Path) -> list:
    records = []
    suffix = file_path.suffix.lower()

    if suffix == ".json":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin1", errors="replace") as f:
                data = json.load(f)

        for row in data:
            raw_disease = row.get("disease", "")
            raw_drug = row.get("potential_repurposed_drug") or row.get("drug_name") or row.get("current_drug", "")
            raw_phase = row.get("clinical_phase") or row.get("status") or "Launched"
            raw_area = row.get("disease_area") or row.get("original_use") or "General"
            raw_moa = row.get("why_repurposed") or row.get("mechanism_of_action") or row.get("mechanism") or ""
            raw_targets = row.get("target_genes") or row.get("targets") or ""
            raw_evidence = row.get("evidence_strength") or ""

            diseases = explode_delimiters(raw_disease)
            if not diseases or not raw_drug:
                continue

            evidence = map_evidence_strength(raw_phase, raw_evidence)
            clean_targets = ", ".join(explode_delimiters(raw_targets))
            moa_str = raw_moa.strip()
            if not moa_str and clean_targets:
                moa_str = f"Modulates {clean_targets}"

            action_type, direction = parse_action_direction(moa_str)

            for dis in diseases:
                dis_title = title_case_disease(dis)
                if not dis_title:
                    continue
                records.append({
                    "disease": dis_title,
                    "drug_name": str(raw_drug).strip(),
                    "clinical_phase": str(raw_phase).strip(),
                    "disease_area": str(raw_area).strip(),
                    "mechanism_of_action": moa_str,
                    "target_genes": clean_targets,
                    "evidence_strength": evidence,
                    "action_type": action_type,
                    "therapeutic_direction": direction,
                })

    elif suffix in (".tsv", ".txt", ".csv"):
        sep = "\t" if suffix in (".tsv", ".txt") else ","
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                valid_lines = [line for line in f if not line.strip().startswith("!")]
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin1", errors="replace") as f:
                valid_lines = [line for line in f if not line.strip().startswith("!")]

        reader = csv.DictReader(valid_lines, delimiter=sep)
        for row in reader:
            raw_disease = row.get("indication") or row.get("disease") or row.get("Disease") or row.get("Indication") or ""
            raw_drug = row.get("pert_iname") or row.get("drug_name") or row.get("drug") or row.get("Drug") or ""
            raw_phase = row.get("clinical_phase") or row.get("phase") or row.get("Phase") or "Launched"
            raw_area = row.get("disease_area") or row.get("area") or row.get("original_use") or "General"
            raw_moa = row.get("moa") or row.get("mechanism_of_action") or row.get("Mechanism") or row.get("why_repurposed") or ""
            raw_targets = row.get("target") or row.get("target_genes") or row.get("targets") or ""
            raw_evidence = row.get("evidence_strength") or ""

            diseases = explode_delimiters(raw_disease)
            if not diseases or not raw_drug:
                continue

            evidence = map_evidence_strength(raw_phase, raw_evidence)
            clean_targets = ", ".join(explode_delimiters(raw_targets))
            moa_str = raw_moa.strip()
            if not moa_str and clean_targets:
                moa_str = f"Modulates {clean_targets}"

            action_type, direction = parse_action_direction(moa_str)

            for dis in diseases:
                dis_title = title_case_disease(dis)
                if not dis_title:
                    continue
                records.append({
                    "disease": dis_title,
                    "drug_name": str(raw_drug).strip(),
                    "clinical_phase": str(raw_phase).strip(),
                    "disease_area": str(raw_area).strip(),
                    "mechanism_of_action": moa_str,
                    "target_genes": clean_targets,
                    "evidence_strength": evidence,
                    "action_type": action_type,
                    "therapeutic_direction": direction,
                })

    elif suffix in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_dict = dict(zip(headers, row))
            raw_disease = str(row_dict.get("disease") or row_dict.get("indication") or row_dict.get("Disease") or "")
            raw_drug = str(row_dict.get("drug_name") or row_dict.get("drug") or row_dict.get("Drug") or "")
            raw_phase = str(row_dict.get("clinical_phase") or row_dict.get("phase") or row_dict.get("Phase") or "Launched")
            raw_area = str(row_dict.get("disease_area") or row_dict.get("area") or "General")
            raw_moa = str(row_dict.get("mechanism_of_action") or row_dict.get("moa") or row_dict.get("Mechanism") or "")
            raw_targets = str(row_dict.get("target_genes") or row_dict.get("targets") or "")
            raw_evidence = str(row_dict.get("evidence_strength") or "")

            diseases = explode_delimiters(raw_disease)
            if not diseases or not raw_drug or raw_drug == "None":
                continue

            evidence = map_evidence_strength(raw_phase, raw_evidence)
            clean_targets = ", ".join(explode_delimiters(raw_targets))
            moa_str = raw_moa.strip()
            if not moa_str and clean_targets:
                moa_str = f"Modulates {clean_targets}"

            action_type, direction = parse_action_direction(moa_str)

            for dis in diseases:
                dis_title = title_case_disease(dis)
                if not dis_title:
                    continue
                records.append({
                    "disease": dis_title,
                    "drug_name": str(raw_drug).strip(),
                    "clinical_phase": str(raw_phase).strip(),
                    "disease_area": str(raw_area).strip(),
                    "mechanism_of_action": moa_str,
                    "target_genes": clean_targets,
                    "evidence_strength": evidence,
                    "action_type": action_type,
                    "therapeutic_direction": direction,
                })

    return records

def seed_database(custom_file=None):
    file_path = find_data_file(custom_file)
    print(f"Reading dataset file: {file_path}")

    records = parse_data_file(file_path)
    print(f"Parsed {len(records):,} exploded repurposing records.")

    # Target SQLite database locations (Canonical DB path + synced locations)
    db_paths = list(dict.fromkeys([
        Path(CANONICAL_DB_PATH),
        BACKEND_DIR / "instance" / "orphaai.db",
        BACKEND_DIR / "orphaai.db",
        PROJECT_ROOT / "instance" / "orphaai.db",
    ]))

    inserted_count = 0
    for db_path in db_paths:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        engine = create_engine(f"sqlite:///{db_path}")
        try:
            RepurposingRecord.__table__.drop(engine, checkfirst=True)
        except Exception:
            pass
        Base.metadata.create_all(engine)

        Session = sessionmaker(bind=engine)
        session = Session()

        for rec in records:
            item = RepurposingRecord(
                disease=rec["disease"],
                drug_name=rec["drug_name"],
                clinical_phase=rec["clinical_phase"],
                disease_area=rec["disease_area"],
                mechanism_of_action=rec["mechanism_of_action"],
                target_genes=rec["target_genes"],
                evidence_strength=rec["evidence_strength"],
                action_type=rec["action_type"],
                therapeutic_direction=rec["therapeutic_direction"],
            )
            session.add(item)

        session.commit()
        inserted_count = session.query(RepurposingRecord).count()
        session.close()
        print(f"Successfully populated database at {db_path} with {inserted_count:,} records.")

    return inserted_count

if __name__ == "__main__":
    custom_path = sys.argv[1] if len(sys.argv) > 1 else None
    count = seed_database(custom_path)
    print(f"Seeding complete! {count:,} repurposing records active in canonical local SQLite database.")
